"""
core/case_manager.py — Advokatlik Ishlar Dosyesi va Mijozlar Bazasi (Case Management).

Mijozlar ma'lumotlari, sud ishlari, tayinlangan sud majlislari sanalari,
da'vo summalari, tayyorlangan yuridik hujjatlar va advokat qaydlarini
SQLite ma'lumotlar bazasida boshqaradi hamda to'liq dosye Word (.docx)
hisobotini shakllantiradi.
"""

from __future__ import annotations
import datetime
import os
import re
import sqlite3
import subprocess
from pathlib import Path
from typing import Dict, Any, List, Optional

import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import sys

def _get_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).parent
        if (exe_dir / "config").is_dir() or (exe_dir / "memory").is_dir():
            return exe_dir
        appdata = os.environ.get("APPDATA")
        if appdata:
            p = Path(appdata) / "AdvokatAI"
            p.mkdir(parents=True, exist_ok=True)
            return p
        return exe_dir
    return Path(__file__).resolve().parent.parent

PROJECT_ROOT = _get_base_dir()
DB_PATH = PROJECT_ROOT / "memory" / "advokat_cases.db"
REPORTS_DIR = (Path(sys.executable).parent / "reports") if getattr(sys, "frozen", False) else (PROJECT_ROOT / "reports")


def get_connection() -> sqlite3.Connection:
    """Ma'lumotlar bazasi ulanishini qaytaradi va jadvallarni initsializatsiya qiladi."""
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row
    cursor = conn.cursor()

    # 1. Mijozlar jadvali
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            phone TEXT,
            email TEXT,
            address TEXT,
            notes TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    """)

    # 2. Sud ishlari jadvali
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS legal_cases (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER,
            case_number TEXT,
            case_title TEXT NOT NULL,
            court_name TEXT,
            judge_name TEXT,
            case_category TEXT DEFAULT 'fuqarolik',
            claim_amount TEXT,
            state_duty TEXT,
            opponent_name TEXT,
            hearing_date TEXT,
            status TEXT DEFAULT 'jarayonda',
            summary TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (client_id) REFERENCES clients (id)
        )
    """)

    # 3. Hujjatlar jadvali (ishga biriktirilgan)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS case_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            case_id INTEGER,
            doc_type TEXT,
            file_name TEXT,
            file_path TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (case_id) REFERENCES legal_cases (id)
        )
    """)

    # 4. Advokatlik qaydlari
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS case_notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            case_id INTEGER,
            note_text TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (case_id) REFERENCES legal_cases (id)
        )
    """)

    conn.commit()
    return conn


def add_or_get_client(name: str, phone: str = "", address: str = "", notes: str = "") -> int:
    """Mijozni bazadan topadi yoki yangi mijoz yaratadi."""
    clean_name = name.strip()
    conn = get_connection()
    c = conn.cursor()
    c.execute("SELECT id FROM clients WHERE LOWER(name) = LOWER(?)", (clean_name,))
    row = c.fetchone()
    if row:
        client_id = row["id"]
    else:
        c.execute("""
            INSERT INTO clients (name, phone, address, notes)
            VALUES (?, ?, ?, ?)
        """, (clean_name, phone.strip(), address.strip(), notes.strip()))
        conn.commit()
        client_id = c.lastrowid
    conn.close()
    return client_id


def create_case(
    client_name: str,
    case_title: str,
    court_name: str = "",
    case_number: str = "",
    judge_name: str = "",
    case_category: str = "fuqarolik",
    claim_amount: str = "",
    state_duty: str = "",
    opponent_name: str = "",
    hearing_date: str = "",
    summary: str = ""
) -> Dict[str, Any]:
    """Yangi sud ishini ochish va bazaga kiritish."""
    client_id = add_or_get_client(client_name)
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Agar ish raqami berilmagan bo'lsa avtomatik hosil qilish
    if not case_number:
        stamp = datetime.datetime.now().strftime("%m%d-%H%M")
        case_number = f"ISH-{stamp}"

    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        INSERT INTO legal_cases (
            client_id, case_number, case_title, court_name, judge_name,
            case_category, claim_amount, state_duty, opponent_name,
            hearing_date, status, summary, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'jarayonda', ?, ?, ?)
    """, (
        client_id, case_number, case_title, court_name, judge_name,
        case_category, claim_amount, state_duty, opponent_name,
        hearing_date, summary, now, now
    ))
    case_id = c.lastrowid
    conn.commit()
    conn.close()

    return {
        "case_id": case_id,
        "case_number": case_number,
        "client_name": client_name,
        "case_title": case_title,
        "court_name": court_name,
        "hearing_date": hearing_date,
        "status": "jarayonda"
    }


def list_active_cases() -> List[Dict[str, Any]]:
    """Hozirgi barcha faol sud ishlari ro'yxatini qaytaradi."""
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        SELECT lc.id, lc.case_number, lc.case_title, lc.court_name, lc.hearing_date,
               lc.status, lc.claim_amount, cl.name as client_name, cl.phone as client_phone
        FROM legal_cases lc
        LEFT JOIN clients cl ON lc.client_id = cl.id
        WHERE lc.status != 'yopildi'
        ORDER BY lc.hearing_date ASC, lc.id DESC
    """)
    rows = c.fetchall()
    res = [dict(r) for r in rows]
    conn.close()
    return res


def find_case(query: str) -> Optional[Dict[str, Any]]:
    """Mijoz ismi, ish raqami yoki sarlavha bo'yicha sud ishini qidirish."""
    conn = get_connection()
    c = conn.cursor()
    q = f"%{query.strip().lower()}%"
    c.execute("""
        SELECT lc.*, cl.name as client_name, cl.phone as client_phone, cl.address as client_address
        FROM legal_cases lc
        LEFT JOIN clients cl ON lc.client_id = cl.id
        WHERE LOWER(lc.case_number) LIKE ?
           OR LOWER(lc.case_title) LIKE ?
           OR LOWER(cl.name) LIKE ?
        ORDER BY lc.id DESC LIMIT 1
    """, (q, q, q))
    row = c.fetchone()
    if not row:
        conn.close()
        return None

    case_data = dict(row)
    case_id = case_data["id"]

    # Hujjatlar
    c.execute("SELECT * FROM case_documents WHERE case_id = ? ORDER BY id DESC", (case_id,))
    case_data["documents"] = [dict(d) for d in c.fetchall()]

    # Qaydlar
    c.execute("SELECT * FROM case_notes WHERE case_id = ? ORDER BY id DESC", (case_id,))
    case_data["notes"] = [dict(n) for n in c.fetchall()]

    conn.close()
    return case_data


def add_case_note(case_id_or_number: str, note_text: str) -> bool:
    """Sud ishiga advokatlik qaydini qo'shish."""
    case = find_case(str(case_id_or_number))
    if not case:
        return False
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        INSERT INTO case_notes (case_id, note_text)
        VALUES (?, ?)
    """, (case["id"], note_text.strip()))
    conn.commit()
    conn.close()
    return True


def update_hearing_date(case_id_or_number: str, new_date: str) -> bool:
    """Sud majlisi sanasi va vaqtini yangilash."""
    case = find_case(str(case_id_or_number))
    if not case:
        return False
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        UPDATE legal_cases
        SET hearing_date = ?, updated_at = ?
        WHERE id = ?
    """, (new_date.strip(), now, case["id"]))
    conn.commit()
    conn.close()
    return True


def close_case(case_id_or_number: str, result_status: str = "yopildi") -> bool:
    """Sud ishini yakunlash."""
    case = find_case(str(case_id_or_number))
    if not case:
        return False
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        UPDATE legal_cases
        SET status = ?, updated_at = ?
        WHERE id = ?
    """, (result_status.strip(), now, case["id"]))
    conn.commit()
    conn.close()
    return True


def attach_document(case_id_or_number: str, doc_type: str, file_path: str) -> bool:
    """Hujjatni sud ishiga biriktirish."""
    case = find_case(str(case_id_or_number))
    if not case:
        return False
    p = Path(file_path)
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        INSERT INTO case_documents (case_id, doc_type, file_name, file_path)
        VALUES (?, ?, ?, ?)
    """, (case["id"], doc_type, p.name, str(p)))
    conn.commit()
    conn.close()
    return True


def build_case_dossier_docx(case: Dict[str, Any], docx_path: Path) -> Path:
    """
    Sud ishi bo'yicha to'liq va rasmiy «ADVOKATLIK DOSYESI» (.docx) hujjatini shakllantiradi.
    """
    doc = docx.Document()

    # Hoshiyalar: Chap 3.0 sm (jildga tikish uchun), O'ng 1.5 sm, Yuqori/Quyi 2.0 sm
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(1.5)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)

    # 1. Bosh Header
    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_after = Pt(2)
    r = p_hdr.add_run("ADVOKAT AI — YURIDIK AMALIYOT VA SUD ISHLARI DOSYESI")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(71, 85, 105)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(8)
    p_line.add_run("─" * 45).font.color.rgb = RGBColor(203, 213, 225)

    # 2. Sarlavha
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(12)
    r_t = p_title.add_run(f"SUD ISHI DOSYESI: № {case.get('case_number', '-')}")
    r_t.bold = True
    r_t.font.size = Pt(15)

    # 3. Ishning pasporti (Jadval shaklida)
    table = doc.add_table(rows=8, cols=2)
    table.style = 'Table Grid'
    for row in table.rows:
        row.height = Pt(22)

    fields = [
        ("Nizo sarlavhasi (Mavzusi):", str(case.get("case_title", "-"))),
        ("Mijoz (Himoya qilinuvchi):", f"{case.get('client_name', '-')} (Tel: {case.get('client_phone') or '-'})"),
        ("Ko'rib chiquvchi sud:", str(case.get("court_name") or "Belgilanmagan")),
        ("Sudya:", str(case.get("judge_name") or "Tayinlanmagan")),
        ("Qarshi taraf (Raqib):", str(case.get("opponent_name") or "-")),
        ("Da'vo bahosi / Davlat boji:", f"{case.get('claim_amount') or '0'} so'm / {case.get('state_duty') or '0'} so'm"),
        ("Sud majlisi sanasi va vaqti:", str(case.get("hearing_date") or "Belgilanmagan")),
        ("Ishning joriy holati:", str(case.get("status", "jarayonda")).upper())
    ]

    for idx, (lbl, val) in enumerate(fields):
        c_left = table.cell(idx, 0)
        c_right = table.cell(idx, 1)
        c_left.width = Cm(5.0)
        c_right.width = Cm(11.5)

        p_l = c_left.paragraphs[0]
        p_l.paragraph_format.space_after = Pt(2)
        r_l = p_l.add_run(lbl)
        r_l.bold = True
        r_l.font.size = Pt(11)

        p_r = c_right.paragraphs[0]
        p_r.paragraph_format.space_after = Pt(2)
        r_r = p_r.add_run(val)
        r_r.font.size = Pt(11)

    # 4. Qisqacha tavsif
    if case.get("summary"):
        p_sum_hdr = doc.add_paragraph()
        p_sum_hdr.paragraph_format.space_before = Pt(14)
        p_sum_hdr.paragraph_format.space_after = Pt(4)
        r_sh = p_sum_hdr.add_run("NIZO MOHIYATI VA ADVOKAT STRATEGIYASI:")
        r_sh.bold = True

        p_sum = doc.add_paragraph()
        p_sum.paragraph_format.first_line_indent = Cm(1.25)
        p_sum.paragraph_format.space_after = Pt(6)
        p_sum.add_run(case["summary"])

    # 5. Ishga biriktirilgan hujjatlar
    p_docs_hdr = doc.add_paragraph()
    p_docs_hdr.paragraph_format.space_before = Pt(12)
    p_docs_hdr.paragraph_format.space_after = Pt(4)
    r_dh = p_docs_hdr.add_run("TAYYORLANGAN VA BIRIKTIRILGAN YURIDIK HUJJATLAR:")
    r_dh.bold = True

    docs = case.get("documents", [])
    if docs:
        for d in docs:
            p_d = doc.add_paragraph()
            p_d.paragraph_format.left_indent = Cm(1.0)
            p_d.paragraph_format.space_after = Pt(2)
            p_d.add_run(f"• [{d.get('doc_type', 'hujjat').upper()}] {d.get('file_name', '-')} ({d.get('created_at', '')})")
    else:
        p_none = doc.add_paragraph()
        p_none.paragraph_format.left_indent = Cm(1.0)
        p_none.add_run("Hozircha rasmiy hujjatlar biriktirilmagan.").italic = True

    # 6. Advokat qaydlari
    notes = case.get("notes", [])
    if notes:
        p_n_hdr = doc.add_paragraph()
        p_n_hdr.paragraph_format.space_before = Pt(12)
        p_n_hdr.paragraph_format.space_after = Pt(4)
        p_n_hdr.add_run("ADVOKATNING KABINET VA SUD QAYDLARI:").bold = True

        for n in notes:
            p_n = doc.add_paragraph()
            p_n.paragraph_format.left_indent = Cm(1.0)
            p_n.paragraph_format.space_after = Pt(2)
            p_n.add_run(f"[{n.get('created_at', '')}] {n.get('note_text', '')}")

    # Imzo
    bugun = datetime.date.today().strftime("%d.%m.%Y")
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(24)
    p_sig.paragraph_format.first_line_indent = Cm(0)
    p_sig.add_run(
        "Advokat / Himoyachi:  ___________________ (ADVOKAT AI)\n"
        f"Dosye chop etilgan sana:  {bugun}\n"
        "M.O'."
    ).bold = True

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    return docx_path


if __name__ == "__main__":
    res = create_case(
        client_name="Ali Valiyev",
        case_title="Aliment undirish va moddiy ta'minot",
        court_name="FIB Shayxontohur tumanlararo sudi",
        case_number="2-1001-2601/77",
        claim_amount="25000000",
        hearing_date="2026-10-24 10:30",
        summary="Javobgardan 2 nafar farzand uchun qonuniy aliment undirish va qarzni to'latish."
    )
    print("Yaratilgan ish:", res)
    active = list_active_cases()
    print("Faol ishlar soni:", len(active))
    dossier = find_case("Ali Valiyev")
    print("Topilgan dosye:", dossier["client_name"], dossier["case_number"])
